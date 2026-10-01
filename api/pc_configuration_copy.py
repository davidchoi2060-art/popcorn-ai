"""Admin-only BOM copy reads; no public recommendation or price authority."""
import hashlib
import json
import re
import unicodedata
from datetime import datetime, timezone
from io import BytesIO
from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import text
from .db import engine
from .part_explanations import is_current
from .taxonomy import SLOT_LABELS
from .pc_catalog_review_snapshot import review_snapshot
from .pc_catalog_changes import load_market, changes_for, issue_groups

router=APIRouter()

QUEUE_LABELS={'ready':'추천 가능 후보','conditional':'조건부 추천 후보','hold':'추천 전 보완','excluded':'우선 제외'}

def queue_status(config, state):
    excluded=(config['status']=='retired' or state['state']=='revoked'
              or any(c['state']=='fail' for c in state['checks'])
              or any('판매중 부품 아님' in b for b in state['blockers']))
    return 'excluded' if excluded else state['recommendation_state']

def review_reasons(state):
    return list(dict.fromkeys(state['blockers']+[c['label']+' · '+c['detail'] for c in state['checks']
        if c['state']=='unknown' and c.get('stage')=='recommendation']))

def queue_counts(rows):
    return {key:sum(r.get('management_state')==key for r in rows) for key in QUEUE_LABELS}

def digest(value):
    return hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':'),default=str).encode()).hexdigest()

def explanation_digest(row):
    return digest({k:row[k] for k in ('content','source_snapshot','source_fingerprint')})

def part_needs_review(part, row, stage='all'):
    if part['pseudo']:
        return False
    if not row or part['explanation_hash']!=explanation_digest(row):
        return True
    if row['content'].get('review_issues'):
        from .pc_review_policy import issue_stages
        if stage != 'recommendation' or any(i['stage']=='recommendation' for i in issue_stages(row)):
            return True
    # Assembly-only identity is linked to the BOM, not to a fictitious retail product.
    if row.get('product_code') is None:
        return row['content'].get('availability_scope')!='assembly_only'
    return not is_current(row)

def read_configuration(conn, identity):
    r=conn.execute(text('SELECT * FROM pc_configurations WHERE configuration_id=:id'),{'id':identity}).mappings().first()
    if r is None: raise HTTPException(404,'조립PC 구성 없음')
    parts=[dict(p) for p in conn.execute(text('SELECT * FROM pc_configuration_parts WHERE configuration_id=:id ORDER BY ordinal'),{'id':identity}).mappings()]
    affected=[]; explanations={}
    for p in parts:
        row=None
        if not p['pseudo']:
            row=conn.execute(text('''SELECT e.*,p.product_name,p.spec_source_text,p.status AS sale_status,p.sale_price
              FROM product_explanations e LEFT JOIN products p ON p.product_code=e.product_code
              WHERE e.source_product_code=:code'''),{'code':p['explanation_code']}).mappings().first()
            p['explanation']=row['content'] if row else {}
            p['image_url']=f"/api/product-images/{p['explanation_code']}/detail"
            p['sale_status']=row['sale_status'] if row else None
            p['current_unit_price']=row['sale_price'] if row else None
            if row: explanations[p['explanation_code']]=row
        p['slot_label']=SLOT_LABELS.get(p['slot'],p['slot'])
        p['needs_review']=part_needs_review(p,row)
        if part_needs_review(p,row): affected.append(p['source_code'])
    offers=[dict(o) for o in conn.execute(text('SELECT * FROM pc_configuration_offers WHERE configuration_id=:id ORDER BY price_snapshot,offer_id'),{'id':identity}).mappings()]
    from .pc_configuration_review import load_review
    current_review=load_review(conn,identity)[3]
    from .pc_media import snapshot
    media=conn.execute(text("SELECT job_id,visual_basis FROM pc_media_jobs WHERE configuration_id=:id AND selected AND status='ready'"),dict(id=identity)).mappings().first()
    media_current=bool(media and snapshot(conn,identity)['visual_basis']==media['visual_basis'])
    workflow=current_review if r['content'].get('_review') else None
    market_alerts=changes_for(r,parts,offers,explanations,*load_market(conn))
    return dict(r,representative_image_url=f"/api/admin/pc-media/{identity}/images/{media['job_id']}" if media_current else None,representative_image_stale=bool(media and not media_current),review_workflow=workflow,parts=parts,offers=offers,needs_review=bool(market_alerts) or (bool(workflow) and not workflow['eligible']) or bool(affected) or bool(r['content'].get('_admin_bom_edit',{}).get('review_required')),affected_parts=affected,
                cooling_plan=current_review['cooling_plan'],market_alerts=market_alerts,current_review=current_review,management_state=queue_status(r,current_review),
                compatibility_display=review_snapshot(identity,parts),
                recommendation_state=current_review['recommendation_state'],assembly_check_count=len(current_review['assembly_checks']),
                customer_publishable=False,price_is_snapshot=True)

def catalog_rows(conn):
    """Bulk read once; evaluate the same source-review rule used by detail reads."""
    rows=conn.execute(text('SELECT * FROM pc_configurations ORDER BY configuration_id')).mappings().all()
    parts=conn.execute(text('SELECT * FROM pc_configuration_parts ORDER BY configuration_id,ordinal')).mappings().all()
    explanations={r['source_product_code']:r for r in conn.execute(text('''
      SELECT e.*,p.product_name,p.spec_source_text,p.status AS sale_status,p.sale_price
      FROM product_explanations e LEFT JOIN products p ON p.product_code=e.product_code
      WHERE EXISTS (SELECT 1 FROM pc_configuration_parts b WHERE b.explanation_code=e.source_product_code)
    ''')).mappings()}
    offers={}
    for o in conn.execute(text('SELECT * FROM pc_configuration_offers ORDER BY price_snapshot,offer_id')).mappings():
        offers.setdefault(o['configuration_id'],[]).append(dict(o))
    by_pc={}
    for p in parts: by_pc.setdefault(p['configuration_id'],[]).append(p)
    from .pc_review_specs import specs_for_review
    from .pc_configuration_review import assess
    by_product={s['product_code']:dict(s) for s in conn.execute(text('SELECT * FROM product_specs WHERE product_code IN (SELECT product_code FROM product_explanations WHERE product_code IS NOT NULL)')).mappings()}
    specs=specs_for_review(explanations,by_product)
    rules=[dict(x) for x in conn.execute(text('SELECT * FROM compat_rules WHERE active ORDER BY rule_id')).mappings()]
    market=load_market(conn)
    from .pc_media import visual_snapshot
    media={r["configuration_id"]:dict(r) for r in conn.execute(text("SELECT configuration_id,job_id,visual_basis FROM pc_media_jobs WHERE selected AND status='ready'")).mappings()}
    result=[]
    for r in rows:
        content=r['content']; bom=by_pc.get(r['configuration_id'],[])
        affected=[p['source_code'] for p in bom if part_needs_review(p,explanations.get(p['explanation_code']))]
        case=next((p for p in bom if p['slot']=='CASE' and not p['pseudo']),None)
        offer_list=offers.get(r['configuration_id'],[])
        current_review=assess(r,bom,offer_list,explanations,specs,rules)
        image=media.get(r['configuration_id'])
        image_current=bool(image and image['visual_basis']==digest(visual_snapshot([p for p in bom if not p['pseudo']],explanations,current_review['cooling_plan'])))
        market_alerts=changes_for(r,bom,offer_list,explanations,*market)
        workflow=current_review if content.get('_review') else None
        offer=offer_list[0] if offer_list else None
        description_ready=bool(content.get('title') and content.get('intro') and bom) and all(
            p['pseudo'] or bool(explanations.get(p['explanation_code'],{}).get('content',{}).get('role')) for p in bom)
        result.append(dict(configuration_id=r['configuration_id'],bom_fingerprint=r['bom_fingerprint'],
            revision=r['revision'],status=r['status'],observed_date=r['observed_date'],
            title=content.get('title',''),intro=content.get('intro',''),facts=content.get('facts',{}),
            search_parts=[{'name':explanations.get(p['explanation_code'],{}).get('product_name')
                          or explanations.get(p['explanation_code'],{}).get('content',{}).get('name',''),
                          'code':p['source_code']} for p in bom if not p['pseudo']],
            source=content.get('source',''),price=offer['price_snapshot'] if offer else None,
            price_note=offer['payload'].get('price_note','') if offer else '',offer_count=len(offer_list),
            image_url=f"/api/admin/pc-media/{r['configuration_id']}/images/{image['job_id']}" if image_current else f"/api/product-images/{case['explanation_code']}/detail" if case else None,
            image_caption='AI 조립 예시 이미지' if image_current else '케이스 이미지',representative_image_stale=bool(image and not image_current),part_count=len(bom),description_ready=description_ready,
            review_state=workflow['state'] if workflow else ('pending' if content.get('_admin_bom_edit',{}).get('review_required') else None),
            cooling_plan=current_review['cooling_plan'],market_alerts=market_alerts,
            management_state=queue_status(r,current_review),review_reasons=review_reasons(current_review),
            recommendation_state=current_review['recommendation_state'],assembly_check_count=len(current_review['assembly_checks']),
            needs_review=bool(market_alerts) or bool(affected) or not description_ready or bool(content.get('_admin_bom_edit',{}).get('review_required')) or (bool(workflow) and not workflow['eligible']),affected_parts=affected,
            customer_publishable=False,price_is_snapshot=True))
    return result


def search_text(value):
    """Ignore model punctuation/spacing, while retaining letters and digits."""
    return re.sub(r'[\W_]+','',unicodedata.normalize('NFKC',str(value)).casefold())


def matches_search(row, query):
    tokens=[search_text(t) for t in query.split() if search_text(t)]
    if not tokens: return True
    facts=row['facts']
    fields=[row['configuration_id'],row['title'],*[str(v) for v in facts.values() if v is not None]]
    for part in row.get('search_parts',[]): fields.extend([part['name'],part['code']])
    # Label installed totals as well as allowing searches by individual part names.
    for key,labels in [('ram_gb',('RAM','메모리')),('storage_gb',('SSD','저장장치'))]:
        value=facts.get(key)
        if isinstance(value,(int,float)):
            fields.extend(f'{label} {value:g}GB' for label in labels)
            if key=='storage_gb' and value>=1000:
                fields.append(f'SSD {value/1000:g}TB')
    normalized=[search_text(f) for f in fields if f]
    return all(any(token in f for f in normalized) for token in tokens)


def filter_rows(rows,q='',source='',review='',visibility='',queue='',changes='',cooling=''):
    return [r for r in rows if
        matches_search(r,q)
        and (not cooling or r.get('cooling_plan',{}).get('method')==cooling)
        and (not source or r['source']==source)
        and (not review or r['needs_review']==(review=='needs_review'))
        and (not queue or r.get('management_state')==queue)
        and (not changes or bool(r.get('market_alerts')))
        and (not visibility or r['customer_publishable']==(visibility=='public'))]


@router.get('/api/admin/pc-configurations')
def list_configurations(offset:int=Query(0,ge=0),limit:int=Query(20,ge=1,le=104),
        q:str=Query('',max_length=100),source:str=Query('',pattern='^(|신규|기존)$'),
        review:str=Query('',pattern='^(|needs_review|ready)$'),visibility:str=Query('',pattern='^(|public|private)$'),
        queue:str=Query('',pattern='^(|ready|conditional|hold|excluded)$'),
        changes:str=Query('',pattern='^(|changed)$'),
        cooling:str=Query('',pattern='^(|air|liquid|bundled_air|unknown)$')):
    with engine.connect() as c: all_rows=catalog_rows(c)
    base=filter_rows(all_rows,q,source,review,visibility,changes=changes,cooling=cooling)
    rows=filter_rows(base,queue=queue)
    return dict(total=len(rows),catalog_total=len(all_rows),offset=offset,limit=limit,
                items=rows[offset:offset+limit],customer_publishable=False,
                summary=queue_counts(base),summary_total=len(base),issue_groups=issue_groups(base),
                changed_count=sum(bool(r.get('market_alerts')) for r in base),checked_at=datetime.now(timezone.utc).isoformat())


def export_workbook(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb=Workbook(); ws=wb.active; ws.title='조립PC 제품군'
    ws.append(['구성 ID','상품명','등록 출처','CPU','GPU','메모리(GB)','저장장치(GB)',
               '기준 가격(원)','가격 기준일','가격 안내','상품 설명','정보 검토','고객 공개','추천 분류','조립 확인 건수','추천 전 보완 사유','가격·판매 확인 알림','냉각 방식','CPU별 냉각 방향'])
    for r in rows:
        f=r['facts']
        values=[r['configuration_id'],r['title'],r['source'],f.get('cpu'),f.get('gpu'),
                f.get('ram_gb'),f.get('storage_gb'),r['price'],str(r['observed_date']),r['price_note'],
                '설명 등록' if r['description_ready'] else '설명 보완 필요',
                '확인 필요' if r['needs_review'] else '정보 변경 없음','비공개',
                QUEUE_LABELS.get(r.get('management_state'),'미확인'),r.get('assembly_check_count',0),' / '.join(r.get('review_reasons',[])),
                ' / '.join(x['label'] for x in r.get('market_alerts',[])),r.get('cooling_plan',{}).get('method_label'),r.get('cooling_plan',{}).get('planning_proposal')]
        ws.append(values)
        # Catalog text is untrusted. Never let a name turn into an Excel formula.
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str): cell.data_type='s'
    ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='186953')
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(60,max(16,max(len(str(c.value or '')) for c in col)+2))
    for row in ws.iter_rows(min_row=2,min_col=8,max_col=8): row[0].number_format='#,##0'
    buffer=BytesIO(); wb.save(buffer); return buffer.getvalue()


@router.get('/api/admin/pc-configurations/export.xlsx')
def export_configurations(q:str=Query('',max_length=100),source:str=Query('',pattern='^(|신규|기존)$'),
        review:str=Query('',pattern='^(|needs_review|ready)$'),visibility:str=Query('',pattern='^(|public|private)$'),
        queue:str=Query('',pattern='^(|ready|conditional|hold|excluded)$'),
        changes:str=Query('',pattern='^(|changed)$'),
        cooling:str=Query('',pattern='^(|air|liquid|bundled_air|unknown)$')):
    with engine.connect() as c: rows=filter_rows(catalog_rows(c),q,source,review,visibility,queue,changes,cooling)
    return Response(export_workbook(rows),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition':'attachment; filename="pc-configurations.xlsx"','Cache-Control':'no-store'})

@router.get('/api/admin/pc-configurations/{identity}')
def get_configuration(identity:str):
    with engine.connect() as c: return read_configuration(c,identity)
