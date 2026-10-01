import unittest
from api.part_explanations import fingerprint
from api.pc_cooling_plan import cooling_plan
from api.pc_configuration_copy import filter_rows, export_workbook
from openpyxl import load_workbook
from io import BytesIO

class CpuCoolingGuidanceTests(unittest.TestCase):
    def plan(self, name, kind='COOLER_CPU_AIR', stale=False):
        parts=[dict(slot='CPU',pseudo=False,explanation_code=1,source_code=1,quantity=1),dict(slot='COOLER',pseudo=False,explanation_code=2,source_code=2,quantity=1)]
        row=dict(product_code=1,product_name=name,spec_source_text='',source_fingerprint=fingerprint(name,''))
        if stale: row['source_fingerprint']='old'
        return cooling_plan(parts,{1:row},{2:dict(part_type=kind)})

    def test_exact_model_suffix(self):
        for model in ('265','265K','265KF','270K Plus','7500X3D'):
            self.assertEqual(self.plan('CPU '+model)['cpu_model'],model)
        self.assertIsNone(self.plan('CPU 1265K')['cpu_model'])

    def test_air_on_official_liquid_recommendation_is_review_not_replacement(self):
        p=self.plan('AMD 9800X3D')
        self.assertEqual(p['method'],'air')
        self.assertIn('수랭 권장',p['manufacturer_guidance'])
        self.assertIn('유지',p['review_hint'])
        self.assertEqual(p['installed'],'separate')
        self.assertFalse(p['verified_bundle'])

    def test_manufacturer_air_and_company_liquid_are_distinct(self):
        self.assertIn('공랭',self.plan('AMD 9600X')['manufacturer_guidance'])
        self.assertIn('수랭 우선',self.plan('Intel 14900KF')['planning_proposal'])
        self.assertNotIn('수랭 권장',self.plan('Intel 14900KF')['manufacturer_guidance'])

    def test_stale_and_unknown_sources_do_not_receive_model_guidance(self):
        self.assertIsNone(self.plan('AMD 9800X3D',stale=True)['cpu_model'])
        self.assertEqual(self.plan('CPU unknown',kind=None)['method'],'unknown')

    def test_actual_liquid_does_not_claim_air(self):
        self.assertEqual(self.plan('CPU 9950X','COOLER_CPU_AIO')['method'],'liquid')

    def test_filters_and_export_keep_cooling_details(self):
        rows=[dict(configuration_id='x',title='CPU 9800X3D',source='신규',facts={},price=1,observed_date=None,price_note='',description_ready=True,needs_review=False,customer_publishable=False,cooling_plan=self.plan('CPU 9800X3D'))]
        self.assertEqual(len(filter_rows(rows,cooling='air')),1)
        self.assertEqual(filter_rows(rows,cooling='liquid'),[])
        sheet=load_workbook(BytesIO(export_workbook(rows))).active
        self.assertEqual(sheet.cell(2,18).value,'별도 공랭')
        self.assertIn('수랭 권장',sheet.cell(2,19).value)
