import unittest
from io import BytesIO
from unittest.mock import patch
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from api import pc_configuration_copy as catalog


class AdminCatalogTests(unittest.TestCase):
    def setUp(self):
        self.rows=[dict(configuration_id='N02',title='문서 PC',source='신규',facts={'cpu':'Core i3-12100','ram_gb':8},
                       needs_review=False,customer_publishable=False,price=509600,observed_date='2026-09-28',
                       price_note='조립비 포함',description_ready=True),
                   dict(configuration_id='P01',title='게임 PC',source='기존',facts={'gpu':'RTX 3050'},
                       needs_review=True,customer_publishable=False,price=1000000,observed_date='2026-09-28',
                       price_note='',description_ready=True)]
        self.rows[0].update(management_state='conditional',assembly_check_count=1,review_reasons=[])
        self.rows[1].update(management_state='hold',assembly_check_count=0,review_reasons=['CPU 소켓 미확인'])
        app=FastAPI();app.include_router(catalog.router);self.client=TestClient(app)

    def test_filter_combinations_and_case_insensitive_specs(self):
        self.assertEqual(catalog.filter_rows(self.rows,q='CORE I3',source='신규',review='ready'),self.rows[:1])
        self.assertEqual(catalog.filter_rows(self.rows,q='RTX',review='needs_review'),self.rows[1:])
        self.assertEqual(catalog.filter_rows(self.rows,source='기존',review='ready'),[])

    def test_info_ready_does_not_mean_public(self):
        self.assertEqual(catalog.filter_rows(self.rows,visibility='public'),[])
        self.assertEqual(catalog.filter_rows(self.rows,visibility='private'),self.rows)

    def test_search_combines_terms_across_parts_in_any_order(self):
        self.rows[0]['search_parts']=[dict(name='GIGABYTE H610M K DDR4',code='12345')]
        self.assertEqual(catalog.filter_rows(self.rows,q='H610M i312100 8GB'),self.rows[:1])
        self.assertEqual(catalog.filter_rows(self.rows,q='8GB 12345 CORE'),self.rows[:1])
        self.assertEqual(catalog.filter_rows(self.rows,q='H610M RTX'),[])
        self.assertEqual(catalog.filter_rows(self.rows,q='ｉ３－１２１００'),self.rows[:1])

    def test_search_capacity_and_empty_query(self):
        self.rows[0]['facts']['storage_gb']=1000
        self.assertEqual(catalog.filter_rows(self.rows,q='RAM8GB SSD1TB'),self.rows[:1])
        self.assertEqual(catalog.filter_rows(self.rows,q='   '),self.rows)
        self.assertEqual(catalog.filter_rows(self.rows,q='없는상품'),[])

    def test_large_pages_and_export_use_same_multi_term_query(self):
        self.rows[0]['search_parts']=[dict(name='화이트 케이스',code='56789')]
        with patch.object(catalog,'catalog_rows',return_value=self.rows):
            result=self.client.get('/api/admin/pc-configurations',params={'q':'케이스 i312100','limit':30}).json()
            r=self.client.get('/api/admin/pc-configurations/export.xlsx',params={'q':'케이스 i312100'})
        self.assertEqual(result['total'],1)
        self.assertEqual(result['items'],self.rows[:1])
        ws=load_workbook(BytesIO(r.content)).active
        self.assertEqual(ws.max_row,2)
        self.assertEqual(ws['A2'].value,result['items'][0]['configuration_id'])

    def test_server_pagination_keeps_filtered_count(self):
        with patch.object(catalog,'catalog_rows',return_value=self.rows):
            r=self.client.get('/api/admin/pc-configurations',params={'offset':1,'limit':1})
        self.assertEqual(r.status_code,200)
        self.assertEqual(r.json()['total'],2);self.assertEqual(r.json()['items'],self.rows[1:])

    def test_invalid_filter_is_not_silently_treated_as_ready(self):
        self.assertEqual(self.client.get('/api/admin/pc-configurations?review=approved').status_code,422)
        self.assertEqual(self.client.get('/api/admin/pc-configurations?queue=approved').status_code,422)

    def test_queue_filter_and_summary_share_search_scope(self):
        with patch.object(catalog,'catalog_rows',return_value=self.rows):
            result=self.client.get('/api/admin/pc-configurations?queue=hold').json()
            narrow=self.client.get('/api/admin/pc-configurations?queue=hold&source=신규').json()
        self.assertEqual(result['items'],self.rows[1:])
        self.assertEqual(result['summary'],dict(ready=0,conditional=1,hold=1,excluded=0))
        self.assertEqual(narrow['items'],[])
        self.assertEqual(narrow['summary_total'],1)
        self.assertEqual(narrow['summary']['conditional'],1)

    def test_exclusion_is_distinct_from_unknown(self):
        state=dict(state='pending',checks=[dict(state='unknown')],blockers=[],recommendation_state='hold')
        self.assertEqual(catalog.queue_status(dict(status='draft'),state),'hold')
        state['checks'][0]['state']='fail'
        self.assertEqual(catalog.queue_status(dict(status='draft'),state),'excluded')
        state['checks']=[];state['state']='revoked'
        self.assertEqual(catalog.queue_status(dict(status='approved'),state),'excluded')

    def test_export_preserves_queue_filter_and_reasons(self):
        with patch.object(catalog,'catalog_rows',return_value=self.rows):
            r=self.client.get('/api/admin/pc-configurations/export.xlsx?queue=hold')
        ws=load_workbook(BytesIO(r.content)).active
        self.assertEqual(ws.max_row,2)
        self.assertEqual(ws['A2'].value,'P01')
        self.assertEqual(ws['N2'].value,'추천 전 보완')
        self.assertEqual(ws['P2'].value,'CPU 소켓 미확인')

    def test_export_route_precedes_identity_and_obeys_filter(self):
        with patch.object(catalog,'catalog_rows',return_value=self.rows):
            r=self.client.get('/api/admin/pc-configurations/export.xlsx?source=신규')
        self.assertEqual(r.status_code,200)
        ws=load_workbook(BytesIO(r.content)).active
        self.assertEqual(ws.max_row,2);self.assertEqual(ws['A2'].value,'N02')
        self.assertEqual(ws['H2'].value,509600)

    def test_export_names_cannot_execute_formulas(self):
        self.rows[0]['title']='=HYPERLINK("https://example.com")'
        ws=load_workbook(BytesIO(catalog.export_workbook(self.rows))).active
        self.assertEqual(ws['B2'].data_type,'s')


if __name__=='__main__':unittest.main()
