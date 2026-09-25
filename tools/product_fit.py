# -*- coding: utf-8 -*-
"""판매 중인 몰 조립PC 를 «용도별 적합성»으로 평가한다 (2026-09-25 재설계 1·2단계).

■ 왜 이 도구가 생겼나
  격자(용도 x 가격 구간)를 칸마다 조합을 «생성»해 채우던 방식을 버리고(사장님
  2026-09-25 「기존거는 모두 무시하고」), 이미 팔고 있고 호환성이 검증된 몰 조립PC
  를 출발점으로 삼는다. 이 도구는 상품마다 「어떤 작업을 어느 수준까지 할 수 있나,
  그 근거는 무엇인가」를 매긴다. 칸을 채우려고 부품을 바꾸지 않는다.

■ 판정 방식
  ① 부품명에서 CPU·GPU 칩을 찾아 성능 지수를 붙인다(CPU_TABLE · GPU_TABLE).
     지수는 공개 벤치마크의 «근사 상대값»이다 — CPU 는 Cinebench R23 멀티/싱글 급,
     GPU 는 RTX 5060 = 100 기준 게임 상대 성능 급. 정밀 측정값이 아니라 등급을
     가르는 용도다. 표에 없는 칩은 지수를 지어내지 않고 「미확인」으로 둔다.
  ② RAM·SSD 용량은 부품명에서 읽는다(「* 2개」 수량 반영).
  ③ 용도마다 기본/쾌적/전문 세 단계의 조건(REQUIREMENTS)을 두고, 조건을 전부
     통과한 가장 높은 단계를 그 상품의 수준으로 적는다. 못 넘은 조건은 근거에 남긴다.

■ 입력: 몰 수집 엑셀(popcornpc-catalog.xlsx — 「상품 목록」·「부품과 기본 옵션」 시트)
■ 출력: 검토용 엑셀(기준표 · 상품별 평가 · 용도 x 가격 미리보기) + JSON
   python3 tools/product_fit.py <catalog.xlsx> <out_dir>
"""
import json
import re
import sys
from collections import defaultdict

# ── CPU 지수: (찾을 패턴, 표시명, 멀티, 싱글, 게임) ─────────────────────────
# 게임 지수 = 게임 프레임에 가까운 등급(싱글 + X3D 캐시 가산). 1.0 = i5-12400 급.
CPU_TABLE = [
    (r"7995WX", "스레드리퍼 7995WX", 150000, 1900, 1.0),
    (r"7975WX", "스레드리퍼 7975WX", 70000, 2000, 1.0),
    (r"7955WX", "스레드리퍼 PRO 7955WX", 38000, 1950, 1.0),
    (r"5975WX", "스레드리퍼 PRO 5975WX", 50000, 1600, 0.8),
    (r"5955WX", "스레드리퍼 PRO 5955WX", 30000, 1600, 0.8),
    (r"3955WX", "스레드리퍼 PRO 3955WX", 25000, 1300, 0.6),
    (r"EPYC 7313P", "EPYC 7313P", 22000, 1300, 0.6),
    (r"9950X3D", "라이젠9 9950X3D", 41000, 2200, 1.55),
    (r"9950X", "라이젠9 9950X", 42000, 2250, 1.35),
    (r"9900X", "라이젠9 9900X", 33000, 2200, 1.3),
    (r"9800X3D", "라이젠7 9800X3D", 23000, 2100, 1.6),
    (r"9700X", "라이젠7 9700X", 21000, 2200, 1.3),
    (r"9600X", "라이젠5 9600X", 17000, 2150, 1.25),
    (r"\b9600\b", "라이젠5 9600", 16500, 2100, 1.2),
    (r"7900X", "라이젠9 7900X", 29000, 2000, 1.2),
    (r"\b7900\b", "라이젠9 7900", 27000, 1950, 1.15),
    (r"7800X3D", "라이젠7 7800X3D", 18000, 1800, 1.5),
    (r"\b7700\b", "라이젠7 7700", 19500, 1950, 1.15),
    (r"7600X", "라이젠5 7600X", 15000, 1950, 1.15),
    (r"7500X3D", "라이젠5 7500X3D", 14000, 1750, 1.35),
    (r"7500F", "라이젠5 7500F", 14500, 1850, 1.1),
    (r"8600G", "라이젠5 8600G", 16000, 1750, 0.95),
    (r"8500G", "라이젠5 8500G", 13000, 1700, 0.9),
    (r"5700X", "라이젠7 5700X", 15000, 1550, 0.9),
    (r"\b5600\b", "라이젠5 5600", 11500, 1550, 0.9),
    (r"5500GT", "라이젠5 5500GT", 11000, 1450, 0.75),
    (r"270K Plus", "코어 울트라7 270K Plus", 40000, 2250, 1.25),
    (r"265KF|265K", "코어 울트라7 265K", 35000, 2200, 1.2),
    (r"울트라7 265", "코어 울트라7 265", 30000, 2150, 1.15),
    (r"250K Plus", "코어 울트라5 250K Plus", 30000, 2200, 1.15),
    (r"245KF|245K", "코어 울트라5 245K", 25000, 2150, 1.1),
    (r"225F|울트라5 225", "코어 울트라5 225", 18000, 2000, 1.0),
    (r"14900K", "i9-14900K", 38000, 2250, 1.35),
    (r"14700K", "i7-14700K", 35000, 2200, 1.3),
    (r"14700", "i7-14700", 30000, 2100, 1.25),
    (r"14600K", "i5-14600K", 24000, 2100, 1.2),
    (r"14400", "i5-14400", 17000, 1850, 1.05),
    (r"14100", "i3-14100", 10500, 1900, 0.95),
    (r"13500", "i5-13500", 21000, 1900, 1.05),
    (r"13400", "i5-13400", 16000, 1800, 1.0),
    (r"12900K", "i9-12900K", 27000, 2000, 1.2),
    (r"12700K", "i7-12700K", 22500, 1950, 1.15),
    (r"12700", "i7-12700", 19500, 1900, 1.1),
    (r"12600K", "i5-12600K", 17500, 1900, 1.1),
    (r"12400", "i5-12400", 12000, 1750, 1.0),
    (r"12100", "i3-12100", 8800, 1700, 0.9),
    (r"프로세서 300\b", "인텔 프로세서 300 (2코어)", 5000, 1650, 0.7),
]

# ── GPU 지수: (패턴, 표시명, 지수, VRAM GB, 종류) ────────────────────────────
# 종류: igpu 내장 · game 지포스/라데온 · pro 워크스테이션 · dc 데이터센터(화면 출력 없음)
GPU_TABLE = [
    (r"H100", "H100 NVL", 0, 94, "dc"),
    (r"PRO 6000", "RTX PRO 6000 Max-Q", 280, 96, "pro"),
    (r"PRO 5000", "RTX PRO 5000", 230, 48, "pro"),
    (r"PRO 4000", "RTX PRO 4000", 150, 24, "pro"),
    (r"PRO 2000", "RTX PRO 2000", 95, 16, "pro"),
    (r"RTX 5090", "RTX 5090", 320, 32, "game"),
    (r"RTX 5080", "RTX 5080", 225, 16, "game"),
    (r"RTX 5070 Ti", "RTX 5070 Ti", 195, 16, "game"),
    (r"RTX 5070", "RTX 5070", 150, 12, "game"),
    (r"RTX 5060 Ti.*16GB", "RTX 5060 Ti 16GB", 120, 16, "game"),
    (r"RTX 5060 Ti", "RTX 5060 Ti 8GB", 118, 8, "game"),
    (r"RTX 5060", "RTX 5060", 100, 8, "game"),
    (r"RTX 3050.*8GB", "RTX 3050 8GB", 60, 8, "game"),
    (r"RTX 3050", "RTX 3050 6GB", 50, 6, "game"),
    (r"RX 9070 XT", "RX 9070 XT", 195, 16, "game"),
    (r"RX 9070", "RX 9070", 175, 16, "game"),
    (r"RX 9060 XT", "RX 9060 XT 16GB", 118, 16, "game"),
    (r"RX 7600", "RX 7600", 85, 8, "game"),
    (r"760M", "Radeon 760M(내장)", 25, 0, "igpu"),
    (r"740M", "Radeon 740M(내장)", 18, 0, "igpu"),
    (r"Radeon 7\b", "Radeon Vega 7(내장)", 10, 0, "igpu"),
    (r"Xe Graphics", "인텔 Xe(내장)", 10, 0, "igpu"),
    (r"UHD", "인텔 UHD(내장)", 6, 0, "igpu"),
    (r"Radeon Graphics", "AMD Radeon(내장 2CU)", 5, 0, "igpu"),
]

LEVELS = ["기본", "쾌적", "전문"]

# ── 기준표: 용도 -> [(단계, 작업 설명, 조건 dict)] ───────────────────────────
# 조건 키: mt 멀티 · st 싱글 · cg CPU게임 · gpu 지수 · vram · ram · ssd(GB)
#          dgpu 외장 그래픽 필요 · nv 엔비디아 필요(CUDA·NVENC)
REQUIREMENTS = {
    "사무용": [
        ("기본", "문서·웹·화상회의 (한글·오피스·크롬 탭 몇 개)", dict(mt=5000, ram=8, ssd=240)),
        ("쾌적", "대용량 엑셀·다중 창·탭 수십 개·듀얼 모니터", dict(mt=12000, ram=16, ssd=480)),
    ],
    "주식·트레이딩": [
        ("기본", "HTS 1~2개 · 모니터 2대", dict(mt=10000, ram=16, ssd=240)),
        ("쾌적", "HTS 여러 개·차트 다수 · 모니터 4대 이상", dict(mt=15000, ram=32, ssd=480, dgpu=1)),
    ],
    "개발": [
        ("기본", "웹·스크립트 개발 (VS Code·브라우저·로컬 서버)", dict(mt=12000, ram=16, ssd=480)),
        ("쾌적", "도커 컨테이너 여러 개 · IDE · 중형 빌드", dict(mt=20000, ram=32, ssd=960)),
        ("전문", "가상머신 다수 · 대형 빌드", dict(mt=30000, ram=64, ssd=960)),
    ],
    "디자인·조판": [
        ("기본", "포토샵·일러스트·인디자인 웹/인쇄물 편집", dict(mt=12000, st=1750, ram=16, ssd=480)),
        ("쾌적", "대형 PSD·다중 레이어·라이트룸 RAW 일괄", dict(mt=16000, st=1850, ram=32, ssd=960, gpu=50, dgpu=1)),
    ],
    "영상편집": [
        ("기본", "프리미어·다빈치 FHD 컷편집 (유튜브)", dict(mt=12000, ram=16, ssd=480, gpu=100, dgpu=1)),
        ("쾌적", "4K 편집·색보정·효과", dict(mt=24000, ram=32, ssd=960, gpu=150, vram=12, dgpu=1)),
        ("전문", "4K 다중 트랙·8K·RAW 영상", dict(mt=33000, ram=64, ssd=1900, gpu=195, vram=16, dgpu=1)),
    ],
    "3D 렌더링": [
        ("기본", "블렌더·C4D 학습·소품 렌더", dict(mt=16000, ram=32, ssd=480, gpu=100, vram=8, dgpu=1)),
        ("쾌적", "실무 장면 렌더 (GPU 렌더러)", dict(mt=24000, ram=32, ssd=960, gpu=195, vram=16, dgpu=1)),
        ("전문", "대형 장면·애니메이션 렌더", dict(mt=30000, ram=64, ssd=1900, gpu=280, vram=24, dgpu=1)),
    ],
    "캐드·설계": [
        ("기본", "오토캐드 2D 도면", dict(st=1700, ram=16, ssd=480)),
        ("쾌적", "솔리드웍스·레빗·인벤터 3D 모델링", dict(st=1950, ram=32, ssd=960, gpu=100, dgpu=1)),
        ("전문", "대형 어셈블리·BIM·구조 해석", dict(mt=30000, st=1950, ram=64, ssd=960, gpu=150, vram=16, dgpu=1)),
    ],
    "음악 작업": [
        ("기본", "큐베이스·FL 작곡·녹음 (트랙 수십 개)", dict(mt=12000, ram=16, ssd=480)),
        ("쾌적", "오케스트라 샘플 라이브러리 · 대형 프로젝트", dict(mt=20000, ram=32, ssd=960)),
    ],
    "방송·스트리밍": [
        ("기본", "FHD 게임 + OBS 송출", dict(mt=12000, ram=16, ssd=480, gpu=100, cg=1.0, nv=1)),
        ("쾌적", "QHD 게임 + 송출 · 녹화 동시", dict(mt=20000, ram=32, ssd=960, gpu=150, cg=1.1, nv=1)),
    ],
    "게임": [
        ("캐주얼", "롤·발로란트·피파·메이플 FHD", dict(gpu=25, ram=16, ssd=240, cg=0.9)),
        ("FHD", "최신 대작 FHD 높음 옵션 60fps 이상", dict(gpu=100, ram=16, ssd=480, cg=1.0)),
        ("QHD", "최신 대작 QHD 높음 옵션", dict(gpu=150, ram=16, ssd=960, cg=1.1)),
        ("4K", "최신 대작 4K 높음 옵션", dict(gpu=225, ram=32, ssd=960, cg=1.2)),
    ],
    "로컬 AI": [
        ("기본", "LLM 7~8B 추론 · 이미지 생성(SDXL)", dict(ram=32, ssd=960, vram=12, nv=1)),
        ("쾌적", "LLM 13~30B · 이미지·영상 생성 본격", dict(ram=64, ssd=960, vram=24, nv=1)),
        ("전문", "LLM 70B급 · 파인튜닝·학습 서버", dict(ram=128, ssd=1900, vram=80, nv=1)),
    ],
}
GAME_LEVELS = ["캐주얼", "FHD", "QHD", "4K"]

COND_KO = {
    "mt": ("CPU 멀티 지수", lambda v: f"{v:,}"), "st": ("CPU 싱글 지수", lambda v: f"{v:,}"),
    "cg": ("CPU 게임 등급", lambda v: f"{v}"), "gpu": ("그래픽 지수", lambda v: f"{v}"),
    "vram": ("그래픽 메모리", lambda v: f"{v}GB"), "ram": ("램", lambda v: f"{v}GB"),
    "ssd": ("SSD", lambda v: f"{v}GB"), "dgpu": ("외장 그래픽", lambda v: "필요"),
    "nv": ("엔비디아 그래픽", lambda v: "필요"),
}


def _qty(name):
    m = re.search(r"\*\s*(\d+)\s*개", name or "")
    return int(m.group(1)) if m else 1


def find_cpu(name):
    for pat, disp, mt, st, cg in CPU_TABLE:
        if re.search(pat, name or ""):
            return dict(name=disp, mt=mt, st=st, cg=cg)
    return None


def find_gpu(name):
    for pat, disp, idx, vram, kind in GPU_TABLE:
        if re.search(pat, name or ""):
            n = _qty(name)
            return dict(name=disp + (f" x{n}" if n > 1 else ""), idx=idx,
                        vram=vram * n, vram_one=vram, kind=kind, n=n)
    return None


def ram_gb(name):
    if not name:
        return None
    m = re.search(r"(\d+)\s*GB", name) or re.search(r"(\d+)G\b", name)
    if not m:
        return None
    return int(m.group(1)) * _qty(name)


def ssd_gb(name):
    if not name:
        return None
    m = re.search(r"(\d+)\s*TB", name)
    if m:
        return int(m.group(1)) * 1000 * _qty(name)
    m = re.search(r"(\d+)\s*GB", name)
    return int(m.group(1)) * _qty(name) if m else None


def specs_of(row):
    cpu = find_cpu(row.get("CPU"))
    gpu = find_gpu(row.get("GPU"))
    # 외장 그래픽이 없으면 CPU 내장 그래픽 이름에서 찾는다
    if gpu is None and cpu is not None:
        gpu = find_gpu(row.get("CPU"))
    return dict(cpu=cpu, gpu=gpu, ram=ram_gb(row.get("RAM")), ssd=ssd_gb(row.get("SSD")))


def check(sp, cond):
    """조건을 못 넘은 항목 목록(비면 통과). 모르는 값은 «미확인»으로 못 넘은 것으로 친다."""
    cpu, gpu = sp["cpu"], sp["gpu"]
    miss = []
    for k, need in cond.items():
        if k in ("mt", "st", "cg"):
            have = cpu and cpu[k]
        elif k == "gpu":
            have = gpu and (gpu["idx"] if gpu["kind"] != "dc" else 0)
        elif k == "vram":
            # AI 는 여러 장의 메모리를 합쳐 쓴다(엔비디아 조건이 붙은 단계) · 그 밖은 한 장 기준
            have = gpu and (gpu["vram"] if cond.get("nv") else gpu["vram_one"])
        elif k == "ram":
            have = sp["ram"]
        elif k == "ssd":
            have = sp["ssd"]
        elif k == "dgpu":
            have = gpu and gpu["kind"] in ("game", "pro")
            if not have:
                miss.append("외장 그래픽 없음")
            continue
        elif k == "nv":
            have = gpu and gpu["kind"] in ("game", "pro", "dc") and not gpu["name"].startswith("RX")
            if not have:
                miss.append("엔비디아 그래픽 아님")
            continue
        if not have:
            miss.append(f"{COND_KO[k][0]} 미확인" if have is None else f"{COND_KO[k][0]} {COND_KO[k][1](need)} 미만")
        elif have < need:
            miss.append(f"{COND_KO[k][0]} {COND_KO[k][1](need)} 미만(현재 {COND_KO[k][1](have)})")
    return miss


def evaluate(sp):
    """용도 -> (도달 단계 또는 None, 다음 단계에서 막힌 이유)."""
    out = {}
    for usage, steps in REQUIREMENTS.items():
        reached, blocked = None, []
        for lvl, _desc, cond in steps:
            miss = check(sp, cond)
            if miss:
                blocked = miss
                break
            reached = lvl
        out[usage] = (reached, blocked)
    return out


def balance_notes(sp, price):
    n = []
    cpu, gpu = sp["cpu"], sp["gpu"]
    if gpu and gpu["kind"] == "game" and gpu["idx"] >= 195 and (sp["ram"] or 0) < 32:
        n.append("고성능 그래픽에 비해 램이 적음")
    if gpu and gpu["kind"] == "game" and gpu["idx"] >= 150 and (sp["ssd"] or 0) < 900:
        n.append("그래픽 급에 비해 SSD 가 작음(대작 게임 몇 개면 참)")
    if cpu and gpu and gpu["kind"] == "game" and gpu["idx"] >= 195 and cpu["cg"] < 1.0:
        n.append("그래픽 급에 비해 CPU 가 약함")
    if cpu and cpu["mt"] >= 30000 and gpu and gpu["kind"] == "igpu" and (sp["ram"] or 0) < 32:
        n.append("CPU 급에 비해 램이 적음")
    return n


def load_catalog(xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    rows = list(wb["상품 목록"].iter_rows(values_only=True))
    h = rows[0]
    items = [dict(zip(h, r)) for r in rows[1:] if r and r[0]]
    return items


def main(xlsx, out_dir):
    import os
    os.makedirs(out_dir, exist_ok=True)
    items = load_catalog(xlsx)
    results = []
    for it in items:
        code = str(it["상품번호"])
        name = (it["상품명"] or "").replace("\n", "")
        price = it["판매가(원)"]
        excl = []
        if code == "123456" or "테스트" in name:
            excl.append("테스트 상품")
        if it.get("사이트 판매 표시") == "품절 표시":
            excl.append("품절")
        if it.get("검토 사항"):
            excl.append(it["검토 사항"])
        sp = specs_of(it)
        if not sp["cpu"] and not excl:
            excl.append("CPU 판독 불가")
        if sp["cpu"] and not excl and (sp["ram"] is None or sp["ssd"] is None):
            excl.append("기본 구성 미선택(램·SSD 없음)")
        ev = evaluate(sp) if sp["cpu"] else {}
        extras = []
        mon = it.get("모니터 기본 선택")
        if mon and "선택하세요" not in mon:
            extras.append("모니터")
        win = it.get("윈도우 기본 선택")
        if win and "선택하세요" not in win:
            extras.append("윈도우")
        results.append(dict(
            code=code, name=name, price=price, member_price=it.get("회원가(원)"),
            route=it.get("발견 경로"), sale=it.get("사이트 판매 표시"),
            cpu_raw=it.get("CPU"), gpu_raw=it.get("GPU"), ram_raw=it.get("RAM"), ssd_raw=it.get("SSD"),
            spec=sp, eval=ev, excluded=excl, includes=extras,
            balance=balance_notes(sp, price) if sp["cpu"] else [],
            url=it.get("원문 URL"),
        ))
    with open(os.path.join(out_dir, "product_fit.json"), "w", encoding="utf-8") as f:
        json.dump(dict(requirements=REQUIREMENTS, results=results), f, ensure_ascii=False, indent=1, default=str)
    write_xlsx(results, os.path.join(out_dir, "판매상품_용도적합성_평가.xlsx"))
    return results


BANDS = [(0, 900000, "~90만"), (900000, 1300000, "90~130만"), (1300000, 1800000, "130~180만"),
         (1800000, 2500000, "180~250만"), (2500000, 3500000, "250~350만"),
         (3500000, 5000000, "350~500만"), (5000000, 10**12, "500만~")]


def band_of(p):
    for lo, hi, lab in BANDS:
        if p is not None and lo <= p < hi:
            return lab
    return "-"


def write_xlsx(results, path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    ws = wb.active
    ws.title = "기준표"
    ws.append(["용도", "단계", "작업 (프로그램·규모)", "조건"])
    for usage, steps in REQUIREMENTS.items():
        for lvl, desc, cond in steps:
            cs = " · ".join(f"{COND_KO[k][0]} {COND_KO[k][1](v)}" + ("" if k in ("dgpu", "nv") else " 이상")
                            for k, v in cond.items())
            ws.append([usage, lvl, desc, cs])
    ws.append([])
    ws.append(["지수 설명", "", "CPU 멀티·싱글 = Cinebench R23 급 근사 · CPU 게임 등급 1.0 = i5-12400 급 · "
               "그래픽 지수 = RTX 5060 = 100 기준 게임 상대 성능 근사 (정밀값 아님, 등급 구분용)"])

    ws2 = wb.create_sheet("상품별 평가")
    usages = list(REQUIREMENTS)
    hdr = ["상품번호", "상품명", "판매가", "가격대", "평가 제외", "가격 포함 품목", "CPU", "그래픽", "램GB", "SSD GB",
           "CPU 멀티", "CPU 싱글", "그래픽 지수", "VRAM"] + usages + ["추천 가능 작업 요약", "다음 단계에서 막힌 이유", "구성 균형 메모", "링크"]
    ws2.append(hdr)
    for r in results:
        sp, ev = r["spec"], r["eval"]
        cpu, gpu = sp["cpu"] or {}, sp["gpu"] or {}
        lv = [(ev.get(u, (None, []))[0] or "-") if ev else "" for u in usages]
        summ = []
        blk = []
        for u in usages:
            if not ev:
                break
            reached, blocked = ev[u]
            if reached:
                desc = next(d for l, d, _ in REQUIREMENTS[u] if l == reached)
                summ.append(f"{u} {reached}: {desc}")
            if blocked and reached != REQUIREMENTS[u][-1][0] and u in ("게임", "영상편집", "3D 렌더링", "로컬 AI", "개발"):
                blk.append(f"{u}: " + ", ".join(blocked[:2]))
        ws2.append([r["code"], r["name"], r["price"], band_of(r["price"]), " · ".join(r["excluded"]),
                    " · ".join(r["includes"]), cpu.get("name") or r["cpu_raw"], gpu.get("name") or r["gpu_raw"],
                    sp["ram"], sp["ssd"], cpu.get("mt"), cpu.get("st"), gpu.get("idx"), gpu.get("vram")]
                   + lv + ["\n".join(summ), "\n".join(blk), " · ".join(r["balance"]), r["url"]])

    ws3 = wb.create_sheet("용도x가격 미리보기")
    ws3.append(["용도·단계"] + [b[2] for b in BANDS])
    live = [r for r in results if not r["excluded"] and r["eval"]]
    for u in usages:
        for lvl, desc, _ in REQUIREMENTS[u]:
            order = [l for l, _, _ in REQUIREMENTS[u]]
            row = [f"{u} {lvl}"]
            for lo, hi, lab in BANDS:
                hits = [r for r in live if r["eval"][u][0] and order.index(r["eval"][u][0]) >= order.index(lvl)
                        and r["price"] is not None and lo <= r["price"] < hi]
                row.append(f"{len(hits)}개 · 최저 {min(h['price'] for h in hits):,}" if hits else "")
            ws3.append(row)

    for w in wb.worksheets:
        for c in w[1]:
            c.font = bold
            c.fill = head_fill
        w.freeze_panes = "B2" if w.title != "기준표" else "A2"
    widths = {"기준표": [14, 8, 44, 90], "용도x가격 미리보기": [22] + [18] * 7}
    for t, ws_ in widths.items():
        for i, wd in enumerate(ws_):
            wb[t].column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wd
    for i, wd in enumerate([9, 48, 11, 11, 18, 12, 22, 22, 7, 8, 9, 9, 9, 7] + [11] * len(usages) + [60, 50, 30, 20]):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wd
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (2, len(hdr) - 3, len(hdr) - 2))
        row[2].number_format = "#,##0"
    wb.save(path)


if __name__ == "__main__":
    res = main(sys.argv[1], sys.argv[2])
    live = [r for r in res if not r["excluded"]]
    print(f"total {len(res)} live {len(live)} excluded {len(res) - len(live)}")
